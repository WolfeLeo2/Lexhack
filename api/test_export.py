"""Checks for POST /api/export (api/export.py): Word and PDF open, carry the database's court words and status (not the
client's), the disclaimer, and refuse oversize bodies.

  uv run python -m api.test_export
"""
import io

import pypdf
from docx import Document
from fastapi.testclient import TestClient

from collections import deque
import time

from . import export, main
from .status import provision_status
from .test_agent import S204

fails = []


def expect(name, got, want):
    if got != want:
        fails.append((name, got, want))


def docx_text(b):
    d = Document(io.BytesIO(b))
    return "\n".join([p.text for p in d.paragraphs] + [p.text for p in d.sections[0].footer.paragraphs])


def pdf_text(b):
    return "\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(b)).pages)


def flat(s):
    return " ".join(s.split())


def run():
    with TestClient(main.app) as c:
        with main.db() as conn:
            st = provision_status(conn, S204)
            ev = next(e for e in st["history"] if e.get("judgment_id"))
            quote = conn.execute("SELECT operative_quote FROM citation_events WHERE event_id = %s",
                                 (ev["event_id"],)).fetchone()[0]
        words = " ".join(quote.split()[:8])
        forged = {"kind": "ruling", "event_id": ev["event_id"], "quote": "The court upheld the section in full.",
                  "case": "Fake v Fake", "citation": "[2099] eKLR", "court": "Nowhere", "paragraph": "1",
                  "checked_by": "a person", "state": None, "url": "https://example.com"}
        section = {"kind": "section", "provision_id": S204, "act": "Penal Code", "number": "204", "heading": None,
                   "status": "in force"}
        parts = [{"kind": "text", "text": "**Section 204** – the court’s “words”:\n\n- first point "}, section,
                 {"kind": "text", "text": "\n"}, forged, {"kind": "text", "text": "Smuggled [[event:1]] here."}]
        main._export_hits.clear()
        for fmt, read in (("docx", docx_text), ("pdf", pdf_text)):
            r = c.post("/api/export", json={"format": fmt, "question": "Is s.204 good law?", "mode": "draft",
                                            "parts": parts, "check": [{"raw_text": "Okuta", "kind": "case",
                                                                       "result": "found", "note": "forged",
                                                                       "flagged": False}]})
            expect(f"{fmt} 200", r.status_code, 200)
            if r.status_code != 200:
                print(r.text)
                continue
            expect(f"{fmt} filename", r.headers["content-disposition"].startswith('attachment; filename="hakiki-draft-'), True)
            text = flat(read(r.content))
            expect(f"{fmt} database quote", flat(words) in text, True)
            expect(f"{fmt} tampered quote gone", "upheld the section in full" in text, False)
            expect(f"{fmt} forged case gone", "Fake v Fake" in text, False)
            expect(f"{fmt} database status", f"[{st['status']}]" in text, True)
            expect(f"{fmt} forged status gone", st["status"] == "in force" or "[in force]" not in text, True)
            expect(f"{fmt} curly quotes and dash", "– the court’s “words”" in text, True)
            expect(f"{fmt} disclaimer", main.DISCLAIMER in text, True)
            expect(f"{fmt} title", "Draft for an advocate's review" in text, True)
            expect(f"{fmt} checked by Hakiki", "Checked by Hakiki" in text, True)
            expect(f"{fmt} forged check note gone", "forged" in text, False)
            expect(f"{fmt} smuggled reference not filled", "Smuggled [[event:1]] here." in text, True)
        r = c.post("/api/export", json={"format": "pdf", "question": "q", "mode": "answer",
                                        "parts": [{"kind": "text", "text": "x" * 210_000}]})
        expect("oversize 413", r.status_code, 413)
        expect("bad format 422", c.post("/api/export", json={"format": "rtf", "question": "q", "parts": []}).status_code, 422)
        a = c.post("/api/export", json={"format": "docx", "question": "q", "mode": "answer", "parts": parts})
        expect("answer title", "Hakiki answer" in docx_text(a.content), True)
        expect("answer has no check", "Checked by Hakiki" in docx_text(a.content), False)
        main._export_hits.clear()
        codes = [c.post("/api/export", json={"format": "rtf", "question": "q", "parts": []}).status_code
                 for _ in range(main.EXPORT_RATE + 1)]
        expect("rate limit", codes[-1], 429)
        main._export_hits.clear()
        test_caps(c)
        test_minors(c, ev["event_id"], words)
    for f in fails:
        print("FAIL", *f)
    print(f"{len(fails)} failures")
    raise SystemExit(bool(fails))


def post(c, parts, question="q", mode="answer", check=None, fmt="docx"):
    main._export_hits.clear()
    return c.post("/api/export", json={"format": fmt, "question": question, "mode": mode, "parts": parts,
                                       "check": check})


def test_caps(c):
    expect("301 parts 422", post(c, [{"kind": "text", "text": "x"}] * 301).status_code, 422)
    expect("30k text 422", post(c, [{"kind": "text", "text": "x" * 30_001}]).status_code, 422)
    held = [main.EXPORT_SLOTS.acquire(blocking=False) for _ in range(2)]
    try:
        expect("slots busy 429", post(c, [{"kind": "text", "text": "x"}]).status_code, 429)
    finally:
        for h in held:
            if h:
                main.EXPORT_SLOTS.release()
    main._export_hits.clear()
    main._export_hits["*"] = deque([time.monotonic()] * main.EXPORT_GLOBAL_RATE)
    r = c.post("/api/export", json={"format": "docx", "question": "q", "parts": []})
    expect("global rate 429", r.status_code, 429)
    main._export_hits.clear()
    expect("slot released after a file", post(c, [{"kind": "text", "text": "x"}]).status_code, 200)


def test_minors(c, eid, words):
    def flags(question, raw):
        r = post(c, [{"kind": "text", "text": "A draft."}], question, "draft",
                 [{"raw_text": raw, "kind": "case", "result": "not_confirmed", "note": "", "flagged": True}])
        return "could not confirm" in docx_text(r.content)
    expect("not_confirmed named in question kept", flags("Draft on Wanjiru v Republic", "Wanjiru v Republic"), True)
    expect("not_confirmed not in question dropped", flags("Draft on s.204", "Wanjiru v Republic"), False)
    expect("not_confirmed with newline dropped", flags("Draft on Wanjiru\nv R", "Wanjiru\nv R"), False)
    expect("not_confirmed non-str dropped", flags("Draft on 5", 5), False)
    split = docx_text(post(c, [{"kind": "text", "text": "[[event:"}, {"kind": "text", "text": f"{eid}]] end"}]).content)
    expect("split reference not re-formed", flat(words) in flat(split), False)
    expect("split reference shown as typed", f"[[event:{eid}]] end" in split, True)
    trail = docx_text(post(c, [{"kind": "text", "text": "see ["}, {"kind": "ruling", "event_id": eid}]).content)
    expect("trailing [ keeps the next ruling", flat(words) in flat(trail), True)
    odd = docx_text(post(c, [{"kind": "ruling", "event_id": {"x": 1}}, {"kind": "section", "provision_id": "a b]]"},
                             {"kind": "case", "judgment_id": True}]).content)
    expect("odd ids removed", odd.count("[reference removed]"), 3)
    expect("odd ids: no stray brackets", "[[" in odd or "]]" in odd, False)
    expect("javascript link dropped", export.web("javascript:alert(1)"), None)
    expect("https link kept", export.web("https://new.kenyalaw.org/x"), "https://new.kenyalaw.org/x")


if __name__ == "__main__":
    run()
