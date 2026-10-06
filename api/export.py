"""POST /api/export: an Ask Hakiki answer or draft as a Word or PDF file. No model call: the file is built from the
parts the client sends, with every ruling (court's words, case, citation, court, paragraph, checked by, state) and
every section's status read again from the database (agent.render_parts), so a file can't carry altered court words
or a forged status. A draft's "Checked by Hakiki" list is re-run on the rebuilt draft (agent.draft_check)."""
import datetime
import io
import logging
import re
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from fpdf import FPDF

logging.getLogger("fontTools").setLevel(logging.WARNING)   # fpdf2 subsets the font and logs every table at INFO

from . import agent

TITLES = {"draft": "Draft for an advocate's review", "answer": "Hakiki answer"}
FONTS = Path(__file__).parent / "fonts"   # Source Serif 4 (OFL): covers ’ “ ” – and the rest of Latin
GREEN = (0x1E, 0x5B, 0x47)                # the site's gazette green, for links
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")   # not allowed in a .docx
BULLET = re.compile(r"^\s*[-*]\s+")
BOLD = re.compile(r"\*\*(.+?)\*\*")


def clean(s):
    return CONTROL.sub("", str(s or ""))


def rebuild(conn, parts):
    """Client parts -> (parts filled in from the database, references that no longer resolve). A reference the
    database doesn't hold becomes 'removed'; prose can't smuggle one in ([[ ]] in text is broken up)."""
    seen, src = {"event": set(), "section": set(), "judgment": set()}, []
    for p in parts:
        k = p.get("kind")
        if k == "text":
            src.append(clean(p.get("text")).replace("[[", "[ [").replace("]]", "] ]"))
            continue
        kind, key = {"ruling": ("event", "event_id"), "section": ("section", "provision_id"),
                     "case": ("judgment", "judgment_id")}.get(k, ("removed", None))
        v = str(p.get(key)) if key else "x"
        seen.setdefault(kind, set()).add(v)
        src.append(f"[[{kind}:{v}]]")
    return agent.render_parts("".join(src), seen, conn)


def checked(conn, parts, removed, question, client_check):
    """The draft's check, re-run on the rebuilt draft. Two warnings depend on the chat run and can't be re-derived
    (a case the user named that Hakiki couldn't confirm; a revision skipped for time): kept from the client, only as
    flags, in the server's words."""
    items = agent.draft_check(conn, parts, question)
    for f in client_check or []:
        if f.get("result") == "not_confirmed" and f.get("raw_text"):
            name = clean(f["raw_text"])[:200]
            items.append(agent.item(name, "case", "not_confirmed",
                                    f"You named {name}. Hakiki could not confirm it, so the draft does not cite it.", True))
        elif f.get("result") == "revision_skipped":
            items.append(agent.item("", "note", "revision_skipped", "Hakiki ran out of time to revise the draft, so "
                                    "this is its first draft with what the check found", True))
    if removed:
        items.append(agent.item("", "note", "references_removed", f"{removed} reference{'s' if removed > 1 else ''} "
                                "couldn't be traced to Hakiki's records and were removed", True))
    return sorted(items, key=lambda i: not i["flagged"])


# ---- layout: parts -> blocks of runs, the same reading as web/lib/answer.ts toLines + Ask's AnswerBody

def locator(p):
    """'18 (see also 20); scope_text from 4' -> 'para 18' (the web's locator, short)."""
    head = re.sub(r"\s*\([^)]*\)", "", (p or "").split(";")[0]).strip()
    return f"para {head}" if head[:1].isdigit() else head


def runs(items):
    """Inline items -> [(text, style, url)]; style '' | 'b' | 'i'."""
    out = []
    for x in items:
        if isinstance(x, str):
            for j, s in enumerate(BOLD.split(x)):
                if s:
                    out.append((s, "b" if j % 2 else "", None))
        elif x["kind"] == "section":
            out.append((f"{x['act']}{' s.' + x['number'] if x['number'] else ''}", "", None))
            out.append((f" [{x['status']}]", "", None))
        elif x["kind"] == "case":
            cite = agent.in_title(x["citation"], x["title"])
            out.append((x["title"], "i", x["url"]))
            if cite:
                out.append((f" {cite}", "", None))
        else:
            out.append(("[reference removed]", "", None))
    return out


def ruling_source(r):
    """The source line under a ruling's words: [(text, style, url)]."""
    out = [(r["case"] or "Parliament (Kenya Law reviser's note)", "i" if r["case"] else "", r["url"])]
    rest = [x for x in (agent.in_title(r["citation"], r["case"]), r["court"], locator(r["paragraph"])) if x]
    by = "Unverified: not yet checked" if r["checked_by"] == "unverified" else r["checked_by"][:1].upper() + r["checked_by"][1:]
    tail = (", " + ", ".join(rest) if rest else "") + f". {by}."
    if r["state"]:
        tail += f" {r['state'][:1].upper()}{r['state'][1:]}."
    return out + [(tail, "", None)]


def blocks(parts):
    """-> [('p' | 'li', runs) | ('ruling', part)]: blank lines split paragraphs, '- ' lines are bullets."""
    lines, cur = [], None
    for p in parts:
        if p["kind"] == "text":
            for k, chunk in enumerate(p["text"].split("\n")):
                if k > 0 or cur is None:
                    m = BULLET.match(chunk)
                    cur = {"bullet": bool(m), "items": []}
                    lines.append(cur)
                    chunk = chunk[m.end():] if m else chunk
                if chunk:
                    cur["items"].append(chunk)
        elif p["kind"] == "ruling":
            lines.append({"ruling": p})
            cur = None
        else:
            if cur is None:
                cur = {"bullet": False, "items": []}
                lines.append(cur)
            cur["items"].append(p)
    out, para = [], []

    def flush():
        if para:
            joined = [x for k, items in enumerate(para) for x in ([" "] if k else []) + items]
            out.append(("p", runs(joined)))
            para.clear()

    for line in lines:
        if "ruling" in line:
            flush()
            out.append(("ruling", line["ruling"]))
        elif all(isinstance(x, str) and not x.strip() for x in line["items"]):
            flush()
        elif line["bullet"]:
            flush()
            out.append(("li", runs(line["items"])))
        else:
            para.append(line["items"])
    flush()
    return out


def document(conn, req, disclaimer):
    """-> (filename, media type, bytes)."""
    parts, removed = rebuild(conn, req.parts)
    check = checked(conn, parts, removed, req.question, req.check) if req.mode == "draft" else None
    today = datetime.datetime.now(ZoneInfo("Africa/Nairobi")).date()
    doc = {"title": TITLES[req.mode], "date": f"{today.day} {today:%B %Y}", "question": clean(req.question),
           "heading": "Draft" if req.mode == "draft" else "Answer", "blocks": blocks(parts), "check": check,
           "disclaimer": disclaimer}
    data = (to_docx if req.format == "docx" else to_pdf)(doc)
    mime = ("application/vnd.openxmlformats-officedocument.wordprocessingml.document" if req.format == "docx"
            else "application/pdf")
    return f"hakiki-{req.mode}-{today.isoformat()}.{req.format}", mime, data


def check_lines(check):
    """'Flagged · Case: X' + the note, one per checked item."""
    kinds = {"case": "Case", "quote": "Quote", "section": "Section", "note": "Note"}
    for f in check:
        raw = f"“{f['raw_text']}”" if f["kind"] == "quote" else f["raw_text"]
        head = f"{'Flagged' if f['flagged'] else 'Checked'} · {kinds.get(f['kind'], f['kind'])}{': ' if raw else ''}"
        yield head, raw, f["note"][:1].upper() + f["note"][1:]


# ---- Word

def add_runs(par, rs):
    for text, style, url in rs:
        if url:
            rid = par.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), rid)
            run = par.add_run(text)
            run.italic, run.font.underline, run.font.color.rgb = style == "i", True, RGBColor(*GREEN)
            link.append(run._r)   # moves the run inside the hyperlink
            par._p.append(link)
        else:
            run = par.add_run(text)
            run.bold, run.italic = style == "b", style == "i"


def to_docx(d):
    w = Document()
    w.core_properties.title, w.core_properties.author = d["title"], "Hakiki"
    w.styles["Normal"].font.name, w.styles["Normal"].font.size = "Georgia", Pt(11.5)
    w.add_heading(d["title"], 0)
    w.add_paragraph(f"{d['date']} · Hakiki (hakiki-ashen.vercel.app)")
    w.add_heading("Question", 2)
    w.add_paragraph(d["question"])
    w.add_heading(d["heading"], 2)
    for kind, x in d["blocks"]:
        if kind == "ruling":
            add_runs(w.add_paragraph(style="Quote"), [(f"“{clean(x['quote'])}”", "", None)])
            src = w.add_paragraph()
            src.paragraph_format.left_indent = Pt(28)
            add_runs(src, [("— ", "", None)] + ruling_source(x))
        else:
            add_runs(w.add_paragraph(style="List Bullet" if kind == "li" else None), x)
    if d["check"] is not None:
        w.add_heading("Checked by Hakiki", 2)
        if not any(f["flagged"] for f in d["check"]):
            w.add_paragraph(f"Hakiki found nothing to flag in what it could check. {len(d['check'])} item"
                            f"{'' if len(d['check']) == 1 else 's'} checked.")
        for head, raw, note in check_lines(d["check"]):
            par = w.add_paragraph(style="List Bullet")
            add_runs(par, [(head, "", None), (raw, "i" if raw else "", None), (f" — {note}", "", None)])
    foot = w.sections[0].footer.paragraphs[0]
    foot.text, foot.alignment = d["disclaimer"], WD_ALIGN_PARAGRAPH.CENTER
    out = io.BytesIO()
    w.save(out)
    return out.getvalue()


# ---- PDF

class Sheet(FPDF):
    disclaimer = ""

    def footer(self):
        self.set_y(-14)
        self.set_font("serif", "", 8.5)
        self.set_text_color(77, 90, 102)
        self.cell(0, 6, f"{self.disclaimer}   Page {self.page_no()}/{{nb}}", align="C")


def to_pdf(d):
    pdf = Sheet(format="A4")
    pdf.disclaimer = d["disclaimer"]
    for style, name in (("", "Regular"), ("B", "Bold"), ("I", "It")):
        pdf.add_font("serif", style, str(FONTS / f"SourceSerif4-{name}.ttf"))
    pdf.set_margins(22, 20, 22)
    pdf.set_auto_page_break(True, 22)
    pdf.set_title(d["title"])
    pdf.set_author("Hakiki")
    pdf.add_page()
    ink = (24, 33, 43)

    def write(rs, size=11.5, h=6.2):
        for text, style, url in rs:
            pdf.set_font("serif", {"b": "B", "i": "I"}.get(style, ""), size)
            pdf.set_text_color(*(GREEN if url else ink))
            pdf.write(h, text, link=url or "")
        pdf.set_text_color(*ink)
        pdf.ln(h)

    def heading(text, size=13, gap=5):
        pdf.ln(gap)
        pdf.set_font("serif", "B", size)
        pdf.set_text_color(*ink)
        pdf.multi_cell(0, size * 0.5, text, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1.5)

    heading(d["title"], 20, 0)
    pdf.set_font("serif", "", 10)
    pdf.set_text_color(77, 90, 102)
    pdf.multi_cell(0, 5, f"{d['date']} · Hakiki (hakiki-ashen.vercel.app)", new_x="LMARGIN", new_y="NEXT")
    heading("Question")
    write([(d["question"], "", None)])
    heading(d["heading"])
    for kind, x in d["blocks"]:
        pdf.ln(1.8)
        if kind == "ruling":
            top, page = pdf.get_y(), pdf.page
            pdf.set_left_margin(30)
            pdf.set_x(30)
            write([(f"“{clean(x['quote'])}”", "", None)], 12)
            pdf.ln(0.8)
            write([("— ", "", None)] + ruling_source(x), 9.5, 5)
            pdf.set_left_margin(22)
            if pdf.page == page:   # a rule down the left of the court's words, as on the site
                pdf.set_draw_color(142, 44, 72)
                pdf.line(26, top, 26, pdf.get_y())
        elif kind == "li":
            pdf.set_font("serif", "", 11.5)
            pdf.write(6.2, "•  ")
            pdf.set_left_margin(27)
            write(x)
            pdf.set_left_margin(22)
        else:
            write(x)
    if d["check"] is not None:
        heading("Checked by Hakiki", gap=7)
        if not any(f["flagged"] for f in d["check"]):
            write([(f"Hakiki found nothing to flag in what it could check. {len(d['check'])} item"
                    f"{'' if len(d['check']) == 1 else 's'} checked.", "", None)], 10.5, 5.5)
        for head, raw, note in check_lines(d["check"]):
            pdf.ln(1.2)
            write([(head, "", None), (raw, "i" if raw else "", None), (f" — {note}", "", None)], 10.5, 5.5)
    return bytes(pdf.output())
